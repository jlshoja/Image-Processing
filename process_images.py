import json
import os
import re
import shutil
import sys
import time
import unicodedata
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl import Workbook
from PIL import Image, ImageFilter, ImageOps
import pillow_heif

pillow_heif.register_heif_opener()

# Persian / Arabic file and folder names crash a Windows console (cp1252/cp1256)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

BASE = Path(__file__).parent
SETTINGS_FILE = Path(os.environ.get("IMAGE_SETTINGS", BASE / "settings.json"))


# --- settings ----------------------------------------------------------------

def _deep_merge(base: dict, override: dict) -> dict:
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(base.get(key), dict):
            _deep_merge(base[key], value)
        else:
            base[key] = value
    return base


def load_settings(path: Path) -> dict:
    """Load settings.json, then merge settings.local.json (machine-specific
    paths, not committed to git) on top of it when it exists."""
    if not path.exists():
        raise SystemExit(f"settings file not found: {path}")
    with path.open(encoding="utf-8") as fh:
        data = json.load(fh)
    local = path.with_name("settings.local.json")
    if local.exists():
        with local.open(encoding="utf-8") as fh:
            _deep_merge(data, json.load(fh))
    for section in (
        "paths",
        "files",
        "encoding",
        "database",
        "color_map",
        "color_rules",
        "color_refs",
        "color_modifiers",
        "color_families",
        "color_space",
    ):
        if section not in data:
            raise SystemExit(f"settings file is missing the '{section}' section: {path}")
    return data


SETTINGS = load_settings(SETTINGS_FILE)
PATHS = SETTINGS["paths"]
FILES = SETTINGS["files"]
ENCODING = SETTINGS["encoding"]
DATABASE = SETTINGS["database"]
COLOR_MAP_CFG = SETTINGS["color_map"]
RULES = SETTINGS["color_rules"]
MODIFIERS = SETTINGS["color_modifiers"]
FAMILIES = SETTINGS["color_families"]
COLOR_DETECTION = SETTINGS.get("color_detection", {})
MATCHING = SETTINGS.get("color_matching", {})
NEUTRAL_CHROMA = MATCHING.get("neutral_chroma", 3.0)
NEUTRAL_SAT = MATCHING.get("neutral_saturation", 0.1)
NEUTRAL_REF_SAT = MATCHING.get("neutral_ref_saturation", 0.15)
LIGHTNESS_WEIGHT = MATCHING.get("lightness_weight", 0.5)
HUE_WEIGHT = MATCHING.get("hue_weight", 0.4)
SATURATION_WEIGHT = MATCHING.get("saturation_weight", 0.1)
NEUTRAL_MISMATCH = MATCHING.get("neutral_mismatch_penalty", 6.0)
IDENTITY_HUE_WEIGHT = MATCHING.get("identity_hue_weight", 0.5)
IDENTITY_SAT_WEIGHT = MATCHING.get("identity_saturation_weight", 0.4)
IDENTITY_LIGHTNESS_WEIGHT = MATCHING.get("identity_lightness_weight", 0.3)
NEUTRAL_PENALTY_PER_CHROMA = MATCHING.get("neutral_penalty_per_chroma", 6.0)
MAX_HUE_PENALTY = MATCHING.get("max_hue_penalty", 90.0)
PRIMARY_COLORS = {c.lower() for c in FAMILIES["primary"]}
FAMILY_MODIFIERS = {m.lower() for m in FAMILIES["modifiers"]}
SYNONYMS = {k.lower(): v.lower() for k, v in FAMILIES.get("synonyms", {}).items()}
COLOR_SPACE = SETTINGS["color_space"]

LANCZOS = Image.Resampling.LANCZOS


def resolve_path(value: str) -> Path:
    path = Path(os.path.expandvars(str(value))).expanduser()
    return path if path.is_absolute() else (BASE / path)


ROOT = resolve_path(PATHS["source_dir"])
CONVERT_ROOT = resolve_path(PATHS["converted_dir"])
PHOTOROOM_ROOT = resolve_path(PATHS.get("photoroom_dir", "Product images photoroom"))
OUTPUT_ROOT = resolve_path(PATHS.get("output_dir", "final image renamed"))
COLOR_REPORT = resolve_path(PATHS["color_report"])
MANIFEST = resolve_path(PATHS["manifest"])
PRODUCT_DB = resolve_path(PATHS["product_db"])
COLOR_MAP = resolve_path(PATHS["color_map"])

MAX_BYTES = ENCODING["max_bytes"]
QUALITY_START = ENCODING["quality_start"]
QUALITY_FLOOR = ENCODING["quality_floor"]
WEBP_METHOD = ENCODING["method"]
CONVERT_MAX_SIDE = ENCODING["convert_max_side"]
BLUR_RADII = ENCODING["blur_radii"]
COLOR_MAX_SIDE = COLOR_DETECTION.get("max_side", 400)
ALPHA_MIN = COLOR_DETECTION.get("alpha_threshold", 128)
BG_TOL = float(COLOR_DETECTION.get("background_tolerance", 12))
CLUSTER_DIST = float(COLOR_DETECTION.get("cluster_distance", 10))
ASSIGN_PENALTY = float(COLOR_DETECTION.get("assignment_penalty", 18))
SUPPORTED = {ext.lower() for ext in FILES["supported_extensions"]}
FEATURED_KEYWORDS = [k.lower() for k in FILES.get("featured_keywords", [])]
FEATURED_SUFFIX = FILES.get("featured_suffix", "main")


def hex_to_rgb(value: str) -> tuple[float, float, float]:
    text = str(value).strip().lstrip("#")
    return (
        int(text[0:2], 16),
        int(text[2:4], 16),
        int(text[4:6], 16),
    )


COLOR_REFS = {
    name: [hex_to_rgb(hexcode) for hexcode in (value if isinstance(value, list) else [value])]
    for name, value in SETTINGS["color_refs"].items()
}


# --- text helpers ------------------------------------------------------------

_PERSIAN_VARIANTS = (
    ("ي", "ی"),
    ("ك", "ک"),
    ("ة", "ه"),
    ("أ", "ا"),
    ("إ", "ا"),
    ("آ", "ا"),
)
_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def norm_text(text) -> str:
    s = str(text)
    for a, b in _PERSIAN_VARIANTS:
        s = s.replace(a, b)
    s = s.replace("\u200c", " ").replace("\u200f", " ").replace("\u200e", " ")
    s = s.replace("-", " ").replace("_", " ").replace("–", " ").replace("—", " ")
    s = re.sub(r"[\"'’]", "", s)
    return re.sub(r"\s+", " ", s).strip().casefold()


def norm_code(value) -> str:
    """Product code as used for matching folder names with database rows.

    Excel often returns integer codes as floats (1234.0) and users mix Persian
    and English digits, so both are normalised.
    """
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).translate(_DIGITS)
    s = re.sub(r"[\u200c\u200e\u200f]", "", s)
    return re.sub(r"\s+", "", s).upper()


def slug(text) -> str:
    s = unicodedata.normalize("NFKD", str(text))  # café -> cafe, not "caf"
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.strip().lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def strip_count(text: str) -> str:
    s = str(text).strip()
    s = re.sub(r"[\s\u200c]*[0-9۰-۹٠-٩]+$", "", s)
    return s.strip()


def is_featured(src: Path) -> bool:
    name = slug(src.stem)
    if not name:
        return False
    tokens = name.split("-")
    return any(keyword in tokens for keyword in FEATURED_KEYWORDS)


# --- color detection ---------------------------------------------------------

def load_for_color(path: Path) -> tuple[Image.Image, np.ndarray | None]:
    """Load an image (downscaled) for color analysis.

    Returns an RGB image composited on white, plus a boolean mask of opaque
    pixels when the file has transparency (None otherwise).
    """
    img = ImageOps.exif_transpose(Image.open(path))
    has_alpha = img.mode in ("RGBA", "LA", "PA") or "transparency" in img.info
    img = img.convert("RGBA" if has_alpha else "RGB")
    img.thumbnail((COLOR_MAX_SIDE, COLOR_MAX_SIDE), LANCZOS)
    if has_alpha:
        alpha = np.asarray(img.getchannel("A")) >= ALPHA_MIN
        white = Image.new("RGBA", img.size, (255, 255, 255, 255))
        return Image.alpha_composite(white, img).convert("RGB"), alpha
    return img, None


def background_rgb(rgb: np.ndarray) -> np.ndarray:
    """Background colour, estimated from a thin ring of the image border.

    Product photos are mostly a flat studio background around the object, so the
    median of the outer border is a solid estimate (and unlike "bright pixels are
    background" it does not erase white or cream products).
    """
    h, w, _ = rgb.shape
    r = max(1, int(min(h, w) * 0.03))
    edge = np.concatenate(
        [
            rgb[:r].reshape(-1, 3),
            rgb[-r:].reshape(-1, 3),
            rgb[:, :r].reshape(-1, 3),
            rgb[:, -r:].reshape(-1, 3),
        ]
    )
    return np.median(edge, axis=0)


def _hsv(rgb) -> tuple[float, float, float]:
    r, g, b = (float(c) / 255.0 for c in rgb)
    mx, mn = max(r, g, b), min(r, g, b)
    delta = mx - mn
    if delta == 0:
        hue = 0.0
    elif mx == r:
        hue = (60 * ((g - b) / delta)) % 360
    elif mx == g:
        hue = 60 * ((b - r) / delta) + 120
    else:
        hue = 60 * ((r - g) / delta) + 240
    sat = 0.0 if mx == 0 else delta / mx
    return float(hue), float(sat), float(mx)


def color_stats(img: Image.Image, alpha: np.ndarray | None = None):
    rgb = np.asarray(img.convert("RGB"), dtype=np.float32)
    hsv = np.asarray(img.convert("HSV"), dtype=np.float32)
    h = hsv[..., 0] * 360.0 / 255.0
    s = hsv[..., 1] / 255.0
    v = hsv[..., 2] / 255.0
    if alpha is not None:
        a = np.asarray(alpha, dtype=bool)
        # a fully opaque "alpha" is just a flattened RGBA file, not a cutout
        keep = a if a.mean() <= 0.98 else None
        if keep is None:
            bg = background_rgb(rgb)
            keep = np.linalg.norm(rgb - bg, axis=2) > BG_TOL
    else:
        # no alpha: the object is everything that differs from the background
        bg = background_rgb(rgb)
        keep = np.linalg.norm(rgb - bg, axis=2) > BG_TOL
    if keep.sum() < RULES["min_pixels"]:
        # the product is (almost) the same tone as the background, e.g. a white
        # bag on a light-grey studio background: treat it as background-coloured
        keep = np.zeros(rgb.shape[:2], dtype=bool)
    return rgb, h, s, v, keep


def object_rgb(img: Image.Image, alpha: np.ndarray | None = None) -> tuple[float, float, float]:
    """Dominant object colour (mode of quantised pixels), not a blend of the
    body with shadows, edges and hardware."""
    rgb, _, _, _, keep = color_stats(img, alpha)
    if keep.sum() == 0:
        return tuple(float(x) for x in background_rgb(rgb))
    px = rgb[keep]
    q = np.clip((px // 16).astype(np.int32), 0, 15)
    keys = q[:, 0] * 256 + q[:, 1] * 16 + q[:, 2]
    counts = np.bincount(keys, minlength=4096)
    sel = px[keys == int(counts.argmax())]
    return tuple(float(x) for x in np.median(sel, axis=0))


def rgb_hex(rgb: tuple[float, float, float]) -> str:
    return "#{:02X}{:02X}{:02X}".format(*(max(0, min(255, int(round(c)))) for c in rgb))


def _name_hsv(h: float, s: float, v: float, v_dark: float) -> str:
    if v_dark < RULES["dark_value"]:
        if s < RULES["black_max_s"]:
            return "black"
        if h >= RULES["hue_red_high"] or h < RULES["hue_brown_high"]:
            return "brown"
        if RULES["hue_brown_high"] <= h < RULES["hue_green_high"]:
            return "green" if s >= RULES["colored_min_s"] else "black"
        if RULES["hue_green_high"] <= h < RULES["hue_teal_high"]:
            return "teal" if s >= RULES["colored_min_s"] else "black"
        if RULES["hue_teal_high"] <= h < RULES["hue_blue_high"]:
            return "navy" if s >= RULES["colored_min_s"] else "black"
        if RULES["hue_blue_high"] <= h < RULES["hue_red_high"]:
            return "purple" if s >= RULES["colored_min_s"] else "black"
    if s < RULES["gray_max_s"]:
        return "white" if v >= RULES["gray_max_v"] else "gray"
    if (h >= RULES["hue_red_high"] or h < RULES["hue_red_low"]) and s > RULES["red_min_s"]:
        return "red"
    if RULES["hue_red_low"] <= h < RULES["hue_brown_high"] and s >= RULES["warm_min_s"]:
        if v < RULES["brown_max_v"]:
            return "brown"
        return "orange" if s > RULES["orange_min_s"] else "beige"
    if RULES["hue_brown_high"] <= h < RULES["hue_yellow_high"]:
        return "yellow"
    if RULES["hue_yellow_high"] <= h < RULES["hue_green_high"]:
        return "green"
    if RULES["hue_green_high"] <= h < RULES["hue_teal_high"]:
        return "teal"
    if RULES["hue_teal_high"] <= h < RULES["hue_blue_high"]:
        return "navy" if v < RULES["navy_max_v"] else "blue"
    if RULES["hue_blue_high"] <= h < RULES["hue_purple_high"]:
        return "purple"
    if RULES["hue_purple_high"] <= h < RULES["hue_pink_high"]:
        return "pink"
    return "red" if s > RULES["red_min_s"] else "brown"


def detect_color(img: Image.Image, alpha: np.ndarray | None = None) -> str:
    _, h, s, v, keep = color_stats(img, alpha)
    if keep.sum() == 0:
        # white/ivory product on a bright background: use the background tone
        bg = background_rgb(np.asarray(img.convert("RGB"), dtype=np.float32))
        hue, sat, val = _hsv(bg)
        if sat < RULES["gray_max_s"]:
            return "white" if val > RULES["gray_max_v"] else "gray"
        return _name_hsv(hue, sat, val, val)
    dom = object_rgb(img, alpha)
    hue, sat, val = _hsv(dom)
    v_dark = float(np.percentile(v[keep], RULES["dark_percentile"]))
    return _name_hsv(hue, sat, val, v_dark)


# --- allowed colors from the product database --------------------------------

def canonical_color(name: str) -> str:
    text = str(name).strip().lower()
    text = SYNONYMS.get(text, text)
    tokens = [SYNONYMS.get(t, t) for t in text.split("-") if t]
    idx = next((i for i, t in enumerate(tokens) if t in PRIMARY_COLORS), None)
    if idx is None:
        idx = next((i for i, t in enumerate(tokens) if t in COLOR_REFS), None)
        if idx is None:
            return text
    prefix = []
    for token in reversed(tokens[:idx]):
        if token in FAMILY_MODIFIERS:
            prefix.insert(0, token)
        else:
            break
    return "-".join(prefix + [tokens[idx]])


def color_ref(name: str) -> list[tuple[float, float, float]] | None:
    """Reference colours for a colour name (a colour may have several shades)."""
    tokens = [t for t in name.split("-") if t]
    # longest run of tokens that is a known color, so "cafe-au-lait" matches as a
    # whole; ties go to the leftmost run
    bases = None
    for size in range(len(tokens), 0, -1):
        for i in range(len(tokens) - size + 1):
            candidate = "-".join(tokens[i:i + size])
            if candidate in COLOR_REFS:
                bases = COLOR_REFS[candidate]
                break
        if bases is not None:
            break
    if bases is None:
        return None
    refs = []
    for base in bases:
        r, g, b = (float(x) for x in base)
        for token, rule in MODIFIERS.items():
            if token in tokens:
                mix = hex_to_rgb(rule["mix"])
                amount = rule["amount"]
                r = r + (mix[0] - r) * amount
                g = g + (mix[1] - g) * amount
                b = b + (mix[2] - b) * amount
        refs.append((r, g, b))
    return refs


def _to_lab(rgb: tuple[float, float, float]) -> np.ndarray:
    c = np.array(rgb, dtype=np.float32) / 255.0
    c = np.where(c <= 0.04045, c / 12.92, ((c + 0.055) / 1.055) ** 2.4)
    m = np.array(COLOR_SPACE["rgb_to_xyz"], dtype=np.float32)
    xyz = m @ c / np.array(COLOR_SPACE["white_point"], dtype=np.float32)
    f = np.where(xyz > 0.008856, np.cbrt(xyz), 7.787 * xyz + 16 / 116)
    return np.array(
        [116 * f[1] - 16, 500 * (f[0] - f[1]), 200 * (f[1] - f[2])], dtype=np.float32
    )


def _lab_chroma_hue(lab: np.ndarray) -> tuple[float, float, float]:
    a, b = float(lab[1]), float(lab[2])
    return float(lab[0]), float(np.hypot(a, b)), float(np.degrees(np.arctan2(b, a)) % 360)


def _hue_diff(h1: float, h2: float) -> float:
    d = abs(h1 - h2) % 360
    return min(d, 360 - d)


def color_distance(
    rgb: tuple[float, float, float],
    ref: tuple[float, float, float] | list[tuple[float, float, float]],
) -> float:
    """Distance between an object colour and a reference colour (or shades).

    Neutral products (white, cream, gray, black) are compared by lightness and
    Lab chroma: a white bag is a light neutral, not a "yellow" one. Coloured
    products are compared by hue first, plus weighted lightness and saturation,
    so a muted sage-green bag is not pulled to a dark brown reference just
    because it is dark. Neutral references get a hue penalty for saturated
    objects, so a navy bag is never matched to black.
    """
    if isinstance(ref, list):
        return min((color_distance(rgb, shade) for shade in ref), default=float("inf"))
    lab = _to_lab(rgb)
    L, C, _ = _lab_chroma_hue(lab)
    ref_lab = _to_lab(ref)
    ref_L, ref_C, _ = _lab_chroma_hue(ref_lab)
    h, s, _ = _hsv(rgb)
    ref_h, ref_s, _ = _hsv(ref)

    if s < NEUTRAL_SAT or C < NEUTRAL_CHROMA:
        base = float(np.hypot(0.5 * (L - ref_L), 0.8 * (C - ref_C)))
        # a gray/white/black bag never *is* a tinted colour: when a neutral
        # reference is available it must win over cream/honey/green references.
        if ref_s >= NEUTRAL_SAT and ref_C >= NEUTRAL_CHROMA:
            base = float(np.hypot(base, NEUTRAL_MISMATCH))
        return base
    if ref_s < NEUTRAL_SAT or ref_C < NEUTRAL_CHROMA:
        hue_deg = MAX_HUE_PENALTY * min(1.0, s / NEUTRAL_REF_SAT)
    else:
        hue_deg = _hue_diff(h, ref_h)
    return float(
        np.sqrt(
            (LIGHTNESS_WEIGHT * (L - ref_L)) ** 2
            + (HUE_WEIGHT * hue_deg) ** 2
            + (SATURATION_WEIGHT * abs(s - ref_s) * 100) ** 2
        )
    )


def nearest_allowed(
    rgb: tuple[float, float, float],
    allowed: list[tuple[str, tuple | None]],
) -> str | None:
    best_name, best_dist = None, None
    for name, ref in allowed:
        if ref is None:
            continue
        dist = color_distance(rgb, ref)
        if best_dist is None or dist < best_dist:
            best_name, best_dist = name, dist
    return best_name


def pick_color(
    img: Image.Image,
    alpha: np.ndarray | None,
    allowed: list[tuple[str, tuple | None]],
) -> tuple[str, str, str]:
    """Return (color to use, detected color name, hex of the object's median color)."""
    detected = detect_color(img, alpha)
    rgb = object_rgb(img, alpha)
    hexcode = rgb_hex(rgb)
    if len(allowed) == 1:
        return allowed[0][0], detected, hexcode
    if allowed:
        match = nearest_allowed(rgb, allowed)
        if match is None:
            # no allowed color has a reference: fall back to the detected name
            exact = [a for a in allowed if a[0] == detected]
            partial = [a for a in allowed if detected in a[0].split("-")]
            match = (exact or partial or [(None, None)])[0][0]
        if match:
            print(
                f"    color: rgb{tuple(int(round(c)) for c in rgb)} "
                f"detected {detected} -> {match}"
            )
            return match, detected, hexcode
    return detected, detected, hexcode


# --- assigning colours to one product's photos -------------------------------

def _identity_distance(rgb1: tuple[float, float, float], rgb2: tuple[float, float, float]) -> float:
    """Whether two photos show the same bag colour.

    Hue and saturation identify a colour; lightness only counts a little because
    shadows and lighting change it between shots of one bag. A neutral bag (white,
    gray, black) is never the same bag as a clearly saturated one.
    """
    L1 = _to_lab(rgb1)[0]
    L2 = _to_lab(rgb2)[0]
    h1, s1, _ = _hsv(rgb1)
    h2, s2, _ = _hsv(rgb2)
    neutral1, neutral2 = s1 < NEUTRAL_SAT, s2 < NEUTRAL_SAT
    if neutral1 and neutral2:
        return float(np.hypot(0.3 * (L1 - L2), 0.6 * abs(s1 - s2) * 100))
    if neutral1 or neutral2:
        return 100.0
    return float(
        np.sqrt(
            (IDENTITY_LIGHTNESS_WEIGHT * (L1 - L2)) ** 2
            + (IDENTITY_HUE_WEIGHT * _hue_diff(h1, h2)) ** 2
            + (IDENTITY_SAT_WEIGHT * abs(s1 - s2) * 100) ** 2
        )
    )


def _cluster_labels(rgbs: list[tuple[float, float, float]], threshold: float) -> list[int]:
    """Greedy clustering of image colours: photos of the same bag share a label."""
    clusters: list[list[int]] = []
    means: list[np.ndarray] = []
    labels: list[int] = []
    for i, rgb in enumerate(rgbs):
        best, best_d = None, None
        for ci, mean in enumerate(means):
            d = _identity_distance(rgb, tuple(float(x) for x in mean))
            if best_d is None or d < best_d:
                best, best_d = ci, d
        if best is not None and best_d <= threshold:
            clusters[best].append(i)
            means[best] = np.mean([rgbs[k] for k in clusters[best]], axis=0)
            labels.append(best)
        else:
            clusters.append([i])
            means.append(np.array(rgb))
            labels.append(len(clusters) - 1)
    return labels


def _assignment(cost: list[list[float]]) -> list[int]:
    """Minimum-cost assignment of rows to distinct columns (rows <= columns).

    Hungarian algorithm; returns the chosen column for every row.
    """
    n, m = len(cost), len(cost[0])
    inf = float("inf")
    u = [0.0] * (n + 1)
    v = [0.0] * (m + 1)
    p = [0] * (m + 1)
    way = [0] * (m + 1)
    for i in range(1, n + 1):
        p[0] = i
        j0 = 0
        minv = [inf] * (m + 1)
        used = [False] * (m + 1)
        while True:
            used[j0] = True
            i0 = p[j0]
            delta = inf
            j1 = 0
            for j in range(1, m + 1):
                if not used[j]:
                    cur = cost[i0 - 1][j - 1] - u[i0] - v[j]
                    if cur < minv[j]:
                        minv[j] = cur
                        way[j] = j0
                    if minv[j] < delta:
                        delta = minv[j]
                        j1 = j
            for j in range(m + 1):
                if used[j]:
                    u[p[j]] += delta
                    v[j] -= delta
                else:
                    minv[j] -= delta
            j0 = j1
            if p[j0] == 0:
                break
        while True:
            j1 = way[j0]
            p[j0] = p[j1]
            j0 = j1
            if j0 == 0:
                break
    result = [-1] * n
    for j in range(1, m + 1):
        if p[j]:
            result[p[j] - 1] = j - 1
    return result


def assign_colors(
    rgbs: list[tuple[float, float, float]],
    allowed: list[tuple[str, tuple | None]],
) -> list[str | None]:
    """Colour name for every image of one product.

    Photos that look alike (several shots of the same bag) are clustered first,
    then clusters are matched to *distinct* colours of the product's list, so two
    different bags never get the same name just because one match was sloppy.
    When there are more clusters than colours, extra clusters reuse the nearest
    colour.
    """
    refs = [(name, ref) for name, ref in allowed if ref is not None]
    if not refs:
        return [None] * len(rgbs)
    if len(refs) == 1:
        return [refs[0][0]] * len(rgbs)
    labels = _cluster_labels(rgbs, CLUSTER_DIST)
    cluster_count = max(labels) + 1
    means = [
        tuple(float(x) for x in np.median([rgbs[i] for i in range(len(rgbs)) if labels[i] == ci], axis=0))
        for ci in range(cluster_count)
    ]
    costs = [[color_distance(mean, ref) for _, ref in refs] for mean in means]
    # one dummy column per cluster: assigning a cluster to a dummy means "this
    # photo set has no colour that fits well enough", so an unused colour is not
    # forced onto a bag (e.g. three photos of one bag and a 6-colour product).
    width = len(refs) + cluster_count
    matrix = [
        row + [ASSIGN_PENALTY] * cluster_count for row in costs
    ]
    picked = _assignment(matrix)
    mapping: dict[int, str] = {}
    for ci, col in enumerate(picked):
        if 0 <= col < len(refs):
            mapping[ci] = refs[col][0]
        else:
            mapping[ci] = refs[min(range(len(refs)), key=lambda j: costs[ci][j])][0]
    return [mapping[label] for label in labels]


# --- input files -------------------------------------------------------------

def _norm_name(text: str) -> str:
    return re.sub(r"[\s_\-]+", "", text).casefold()


def locate_file(path: Path) -> tuple[Path | None, list[Path]]:
    """Find an input workbook, tolerating small differences in the file name.

    Returns (found file or None, Excel files in the same folder).
    """
    if path.is_file():
        return path, []
    folder = path.parent
    if not folder.is_dir():
        return None, []
    candidates = [
        p for p in sorted(folder.glob("*.xls*")) if not p.name.startswith("~$")
    ]
    want = _norm_name(path.stem)
    exact = [p for p in candidates if _norm_name(p.stem) == want]
    if len(exact) == 1:
        return exact[0], candidates
    close = [p for p in candidates if _norm_name(p.stem).startswith(want)]
    if len(close) == 1:
        return close[0], candidates
    return None, candidates


def require_inputs() -> None:
    """Stop before touching any image if the database or color map is missing."""
    global PRODUCT_DB, COLOR_MAP
    for label, name in (("product database", "PRODUCT_DB"), ("color mapping", "COLOR_MAP")):
        path = globals()[name]
        found, candidates = locate_file(path)
        if found is None:
            lines = [f"{label} not found: {path}"]
            if candidates:
                lines.append("Excel files in that folder:")
                lines += [f"  {p.name}" for p in candidates]
            lines.append("Fix the path in settings.local.json and run again. Nothing was processed.")
            raise SystemExit("\n".join(lines))
        if found != path:
            print(f"{label}: using '{found.name}' (configured: '{path.name}')")
            globals()[name] = found


# --- database ----------------------------------------------------------------

def _sheet(workbook, cfg: dict):
    name = cfg.get("sheet")
    if name in (None, ""):
        return workbook[workbook.sheetnames[0]]
    return workbook[name]


def _column_index(header_values, target: str | None, fallback: int) -> int:
    if target:
        wanted = norm_text(target)
        for i, value in enumerate(header_values or ()):
            if value and norm_text(value) == wanted:
                return i
    return fallback - 1


def load_color_map() -> dict[str, str]:
    if not COLOR_MAP.exists():
        print(f"warning: color mapping file not found: {COLOR_MAP}")
        return {}
    wb = openpyxl.load_workbook(COLOR_MAP, read_only=True, data_only=True)
    try:
        ws = _sheet(wb, COLOR_MAP_CFG)
        header_row = COLOR_MAP_CFG["header_row"]
        header = list(
            ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        persian_col = _column_index(
            header[0] if header else (),
            COLOR_MAP_CFG.get("persian_header"),
            COLOR_MAP_CFG["persian_column"],
        )
        english_col = _column_index(
            header[0] if header else (),
            COLOR_MAP_CFG.get("english_header"),
            COLOR_MAP_CFG["english_column"],
        )
        mapping = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            persian = row[persian_col] if len(row) > persian_col else None
            english = row[english_col] if len(row) > english_col else None
            if persian and english:
                mapping.setdefault(norm_text(persian), slug(english))
        return mapping
    finally:
        wb.close()


def resolve_color(raw, color_map: dict[str, str]) -> str | None:
    text = norm_text(raw)
    if not text:
        return None
    tries = [text]
    stripped = norm_text(strip_count(raw))
    if stripped and stripped != text:
        tries.append(stripped)
    variants = []
    for t in tries:
        variants.append(t)
        words = t.split()
        if words and len(words[-1]) > 2 and words[-1].endswith("ی"):
            variants.append(" ".join(words[:-1] + [words[-1][:-1]]))
        if t.endswith("سورمه"):
            variants.append(t + " ای")
        if "سورمه" in t:
            variants.append(t.replace("سورمه", "سرمه"))
        if t.endswith("نسکافه"):
            variants.append(t + " ای")
    for v in variants:
        if v in color_map:
            return color_map[v]
    # fallback: longest mapped color name that appears as whole word(s) in the text
    padded = f" {text} "
    best = None
    for key, english in color_map.items():
        if key and f" {key} " in padded and (best is None or len(key) > len(best[0])):
            best = (key, english)
    return best[1] if best else None


def _product_rows():
    if not PRODUCT_DB.exists():
        print(f"warning: product database not found: {PRODUCT_DB}")
        return
    wb = openpyxl.load_workbook(PRODUCT_DB, read_only=True, data_only=True)
    try:
        ws = _sheet(wb, DATABASE)
        header_row = DATABASE["header_row"]
        header = list(
            ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        code_col = _column_index(
            header[0] if header else (),
            DATABASE.get("code_header"),
            DATABASE["code_column"],
        )
        color_col = _column_index(
            header[0] if header else (),
            DATABASE.get("color_header"),
            DATABASE["color_column"],
        )
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            code = row[code_col] if len(row) > code_col else None
            raw = row[color_col] if len(row) > color_col else None
            if code not in (None, ""):
                yield norm_code(code), "" if raw is None else str(raw)
    finally:
        wb.close()


def load_database() -> tuple[dict[str, list[tuple[str, tuple | None]]], list[tuple[str, str, str]]]:
    """Read the product database and color mapping once.

    Returns (products, missing):
    - products: product code -> allowed colors (English), e.g. 'آبی ۱ | قرمز ۲'
      becomes [blue, red]. Codes with an empty color cell map to [].
    - missing: (code, color text, problem) for colors without a mapping entry.
    """
    color_map = load_color_map()
    products: dict[str, list[tuple[str, tuple | None]]] = {}
    missing: list[tuple[str, str, str]] = []
    seen = set()
    for code, raw in _product_rows():
        allowed = products.setdefault(code, [])
        for segment in re.split(r"[|,،;؛\n]", raw):
            if not strip_count(segment).strip():
                continue  # empty, or just a count
            english = resolve_color(segment, color_map)
            if not english:
                key = (code, norm_text(segment))
                if key not in seen:
                    seen.add(key)
                    missing.append((code, segment.strip(), "no entry in color mapping"))
                continue
            english = canonical_color(english)
            if english not in {name for name, _ in allowed}:
                allowed.append((english, color_ref(english)))
    # In a multi-color product a color without a reference can never be chosen
    # by the nearest-color match, so tell the user to add it to color_refs.
    for code, allowed in products.items():
        if len(allowed) > 1:
            for name, ref in allowed:
                if ref is None:
                    missing.append(
                        (code, name, "no reference color: add it to color_refs in settings.json")
                    )
    return products, missing


# --- option 1: convert to webp under the size limit --------------------------

def prepare(img: Image.Image) -> Image.Image:
    img = ImageOps.exif_transpose(img)  # phone photos: apply the rotation flag
    if img.mode in ("RGBA", "LA", "PA", "P"):
        img = img.convert("RGBA")
        white = Image.new("RGBA", img.size, (255, 255, 255, 255))
        img = Image.alpha_composite(white, img)
    return img.convert("RGB")


def compress(
    img: Image.Image,
    dest: Path,
    max_bytes: int = MAX_BYTES,
    allow_blur: bool = True,
) -> int:
    """Save img as WebP at dest, as small as needed to fit max_bytes.

    The file left on disk is always the one for the returned quality.
    """
    def encode_at(im: Image.Image, quality: int) -> int:
        im.save(dest, format="WEBP", quality=quality, method=WEBP_METHOD)
        return dest.stat().st_size

    # one save is enough in the common case: start quality already fits
    if encode_at(img, QUALITY_START) <= max_bytes:
        return QUALITY_START

    def search(im: Image.Image) -> int | None:
        lo, hi = QUALITY_FLOOR, QUALITY_START - 1
        best, last = None, None
        while lo <= hi:
            mid = (lo + hi) // 2
            last = mid
            if encode_at(im, mid) <= max_bytes:
                best = mid
                lo = mid + 1
            else:
                hi = mid - 1
        if best is not None and best != last:
            encode_at(im, best)  # the last probe may have been a larger file
        return best

    quality = search(img)
    if quality is not None:
        return quality
    last_im = img
    if allow_blur:
        for radius in BLUR_RADII:
            last_im = img.filter(ImageFilter.GaussianBlur(radius))
            quality = search(last_im)
            if quality is not None:
                return quality
    size = encode_at(last_im, QUALITY_FLOOR)
    print(
        f"    warning: {dest.name} is still {size / 1024:.1f} KB "
        f"(limit {max_bytes / 1024:.0f} KB) at quality {QUALITY_FLOOR}"
    )
    return QUALITY_FLOOR


def convert_file(src: Path, dest: Path) -> tuple[int, str, float]:
    started = time.perf_counter()
    if src.suffix.lower() == ".webp" and src.stat().st_size <= MAX_BYTES:
        shutil.copyfile(src, dest)
        return QUALITY_START, "copied", time.perf_counter() - started
    img = prepare(Image.open(src))
    if CONVERT_MAX_SIDE and max(img.size) > CONVERT_MAX_SIDE:
        img.thumbnail((CONVERT_MAX_SIDE, CONVERT_MAX_SIDE), LANCZOS)
    quality = compress(img, dest, MAX_BYTES)
    return quality, "compressed", time.perf_counter() - started


def find_sources() -> list[Path]:
    return [
        p
        for p in sorted(ROOT.rglob("*"))
        if p.is_file()
        and p.suffix.lower() in SUPPORTED
        and not set(p.relative_to(ROOT).parts).intersection(
            {CONVERT_ROOT.name, OUTPUT_ROOT.name}
        )
    ]


def convert(sources: list[Path]) -> list[Path]:
    converted = []
    planned: dict[Path, Path] = {}
    total = len(sources)
    print(f"found {total} source file(s) in {ROOT}  ->  {CONVERT_ROOT}", flush=True)
    for i, src in enumerate(sources, 1):
        pct = i / total * 100 if total else 100
        rel = src.relative_to(ROOT)
        dest = (CONVERT_ROOT / rel.parent / f"{src.stem}.webp").resolve()
        if dest in planned and planned[dest] != src:
            # a.jpg and a.png in one folder would both become a.webp
            alt = dest.with_name(f"{src.stem}_{src.suffix.lstrip('.').lower()}.webp")
            print(f"    warning: {rel} would overwrite {planned[dest].name}, using {alt.name}")
            dest = alt
        planned[dest] = src
        dest.parent.mkdir(parents=True, exist_ok=True)
        if dest.exists():
            converted.append(dest)
            continue
        # write to a temp file so an interrupted run never leaves a broken .webp
        tmp = dest.with_name(f"{dest.stem}.part.webp")
        try:
            quality, mode, seconds = convert_file(src, tmp)
            tmp.replace(dest)
        finally:
            if tmp.exists():
                tmp.unlink()
        converted.append(dest)
        kb = dest.stat().st_size / 1024
        print(
            f"[{i}/{total} {pct:.0f}%] convert: {rel} -> "
            f"{CONVERT_ROOT.name}/{dest.relative_to(CONVERT_ROOT)}  {kb:.1f} KB  "
            f"q={quality}  {mode}  {seconds:.1f} s"
        )
    print(f"[100%] convert: {total} file(s) checked.\n")
    return converted


# --- manifest & report -------------------------------------------------------

def load_manifest() -> dict:
    path = OUTPUT_ROOT / MANIFEST.name
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}


def save_manifest(manifest: dict) -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_ROOT / MANIFEST.name
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)


def source_key(src: Path) -> str:
    return src.relative_to(PHOTOROOM_ROOT).as_posix()


def write_color_report(records: list[dict], missing: list[tuple[str, str, str]]) -> bool:
    wb = Workbook()
    ws = wb.active
    ws.title = "detected"
    ws.append(
        [
            "source",
            "product code",
            "output file",
            "color (used)",
            "detected",
            "hex",
            "featured",
            "from product list",
            "allowed colors",
        ]
    )
    for rec in records:
        ws.append(
            [
                rec.get("source", ""),
                rec.get("code", ""),
                rec.get("output", ""),
                rec.get("color", ""),
                rec.get("detected", ""),
                rec.get("hex", ""),
                "yes" if rec.get("featured") else "no",
                "yes" if rec.get("matched") else "no",
                rec.get("allowed", ""),
            ]
        )
    ws2 = wb.create_sheet("missing")
    ws2.append(["product code", "color", "problem"])
    for code, raw, problem in missing:
        ws2.append([code, raw, problem])
    for sheet in (ws, ws2):
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 60)
    COLOR_REPORT.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(COLOR_REPORT)
    except PermissionError:
        print(
            f"warning: could not write {COLOR_REPORT.name} "
            "(close it in Excel and run again to update it)"
        )
        return False
    return True


# --- option 2: rename photoroom images by color ------------------------------

def find_photoroom_sources() -> list[Path]:
    if not PHOTOROOM_ROOT.exists():
        raise SystemExit(f"photoroom folder not found: {PHOTOROOM_ROOT}")
    return [
        p
        for p in sorted(PHOTOROOM_ROOT.rglob("*"))
        if p.is_file() and p.suffix.lower() in SUPPORTED
    ]


def _glob_escape(text: str) -> str:
    return re.sub(r"([\[\]*?])", r"[\1]", text)


def _name_taken(out: str, folder: str, existing_stems: set[str]) -> bool:
    if out in existing_stems:
        return True
    target = OUTPUT_ROOT / folder
    return target.exists() and any(target.glob(f"{_glob_escape(out)}.*"))


def _records(manifest: dict) -> list[dict]:
    return [
        dict(entry, source=key)
        for key, entry in manifest.items()
        if isinstance(entry, dict) and entry.get("output")
    ]


def process(sources: list[Path]) -> None:
    require_inputs()
    manifest = load_manifest()

    # an output that was deleted by hand must be produced again
    stale = [
        key
        for key, entry in manifest.items()
        if not (isinstance(entry, dict) and entry.get("output")
                and (OUTPUT_ROOT / entry["output"]).exists())
    ]
    for key in stale:
        manifest.pop(key)
    if stale:
        print(f"{len(stale)} output file(s) missing on disk, they will be processed again")

    existing_stems = {Path(entry["output"]).stem for entry in manifest.values()}
    pending = [src for src in sorted(sources) if source_key(src) not in manifest]
    print(
        f"found {len(sources)} image(s) in {PHOTOROOM_ROOT.name}, "
        f"{len(pending)} new ({len(sources) - len(pending)} already in manifest)",
        flush=True,
    )

    print("reading product database and color mapping...", flush=True)
    products, missing = load_database()
    no_ref = sorted({raw for _, raw, problem in missing if problem.startswith("no reference color")})
    if no_ref:
        print(
            "warning: these colors have no reference color and cannot be matched "
            f"(add them to color_refs in settings.json): {', '.join(no_ref)}"
        )

    if not pending:
        print("process: no new images.")
        records = _records(manifest)
        if records:
            write_color_report(records, missing)
            print(
                f"color report: {COLOR_REPORT.name} "
                f"({len(records)} image(s), {len(missing)} missing color(s))"
            )
        return

    total = len(pending)
    done = 0
    announced = set()
    no_colors = []
    failed = []

    groups: dict[Path, list[Path]] = {}
    for src in sorted(pending):
        groups.setdefault(src.parent, []).append(src)

    for folder, files in groups.items():
        folder_name = folder.name
        if folder == PHOTOROOM_ROOT:
            for src in files:
                done += 1
                print(f"[{done}/{total}] skip (not inside a product-code folder): {src.name}")
            continue
        code = norm_code(folder_name)
        allowed = products.get(code, [])

        if folder_name not in announced:
            announced.add(folder_name)
            if allowed:
                names = ", ".join(name for name, _ in allowed)
                print(f"colors for {folder_name} (from product list): {names}")
            else:
                reason = (
                    "color list is empty"
                    if code in products
                    else "code not found in database"
                )
                no_colors.append((folder_name, "", reason))
                print(f"colors for {folder_name}: {reason}, using detected color")

        if len(allowed) == 1:
            # single-color product: no detection needed
            entries = [(src, None, "single color", "") for src in files]
        else:
            entries = []
            for src in files:
                try:
                    img, alpha = load_for_color(src)
                    rgb = object_rgb(img, alpha)
                    entries.append((src, rgb, detect_color(img, alpha), rgb_hex(rgb)))
                except Exception as exc:  # unreadable/corrupt image: report and go on
                    done += 1
                    print(f"[{done}/{total}] ERROR {source_key(src)}: {exc}")
                    failed.append((folder_name, src.name, f"could not read image: {exc}"))

        if len(allowed) == 1:
            colors = [allowed[0][0]] * len(entries)
        else:
            colors = assign_colors([entry[1] for entry in entries], allowed)

        for (src, rgb, detected, hexcode), assigned in zip(entries, colors):
            done += 1
            color = assigned or detected
            ref_by_name = {name: ref for name, ref in allowed}
            distance = (
                color_distance(rgb, ref_by_name[color])
                if rgb is not None and color in ref_by_name
                else None
            )
            uncertain = distance is not None and distance > ASSIGN_PENALTY
            if rgb is not None:
                print(
                    f"    color: rgb{tuple(int(round(c)) for c in rgb)} "
                    f"detected {detected} -> {color}"
                    + (f"  [check: distance {distance:.1f}]" if uncertain else "")
                )

            featured = is_featured(src)
            seq = 1
            if featured:
                out = f"{folder_name}-{color}-{FEATURED_SUFFIX}"
                while _name_taken(out, folder_name, existing_stems):
                    out = f"{folder_name}-{color}-{FEATURED_SUFFIX}-{seq}"
                    seq += 1
            else:
                out = f"{folder_name}-{color}-{seq:02d}"
                while _name_taken(out, folder_name, existing_stems):
                    seq += 1
                    out = f"{folder_name}-{color}-{seq:02d}"

            dest = OUTPUT_ROOT / folder_name / f"{out}{src.suffix.lower()}"
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest)  # no conversion: same file, new name

            key = source_key(src)
            manifest[key] = {
                "output": dest.relative_to(OUTPUT_ROOT).as_posix(),
                "code": folder_name,
                "color": color,
                "detected": detected,
                "hex": hexcode,
                "featured": featured,
                "matched": color in {name for name, _ in allowed},
                "distance": None if distance is None else round(distance, 2),
                "uncertain": uncertain,
                "allowed": ", ".join(name for name, _ in allowed),
            }
            existing_stems.add(dest.stem)
            save_manifest(manifest)  # save progress after every image
            print(
                f"[{done}/{total} {done / total * 100:.0f}%] {key} -> "
                f"{OUTPUT_ROOT.name}/{dest.relative_to(OUTPUT_ROOT)}"
            )

    save_manifest(manifest)
    missing = missing + no_colors + failed
    records = _records(manifest)
    write_color_report(records, missing)
    print(
        f"color report: {COLOR_REPORT.name} "
        f"({len(records)} image(s), {len(missing)} problem(s))"
    )
    print("[100%] process: all images done.")


# --- runner ------------------------------------------------------------------

def main() -> None:
    mode = sys.argv[1].lower() if len(sys.argv) > 1 else ""
    if mode not in ("convert", "process"):
        print("Usage: python process_images.py [convert|process]")
        print(f"  settings: {SETTINGS_FILE}")
        print(f"  convert - option 1: convert all sources to webp under {MAX_BYTES // 1024} KB, keeping original names/folders")
        print(f"  process - option 2: read '{PHOTOROOM_ROOT.name}', name each image by its")
        print(f"            color (matched with the product database) and save to '{OUTPUT_ROOT.name}/<product code>/'")
        return
    print(f"settings: {SETTINGS_FILE}", flush=True)
    if mode == "convert":
        n = len(convert(find_sources()))
        print(f"\n{n} file(s) ensured as webp in {CONVERT_ROOT.name}/.")
    else:
        process(find_photoroom_sources())
        print(f"\nDone: outputs in {OUTPUT_ROOT.name}/<product code>/.")


if __name__ == "__main__":
    main()
