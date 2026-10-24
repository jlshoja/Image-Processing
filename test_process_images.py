from pathlib import Path

import numpy as np
import openpyxl
import pytest
from PIL import Image

import process_images as pi


def solid(color, size=(120, 120)):
    return Image.new("RGB", size, color)


def on_white(color, size=200, box=100):
    img = Image.new("RGB", (size, size), (255, 255, 255))
    off = (size - box) // 2
    img.paste(Image.new("RGB", (box, box), color), (off, off))
    return img


def ref(name):
    return pi.color_ref(name)


# --- text helpers ------------------------------------------------------------

def test_norm_text_unifies_arabic_variants_and_zwnj():
    assert pi.norm_text("كرمي") == pi.norm_text("کرمی")
    assert pi.norm_text("نسکافه\u200cای") == "نسکافه ای"


def test_strip_count_removes_trailing_numbers():
    assert pi.strip_count("آبی ۱") == "آبی"
    assert pi.strip_count("قرمز 2") == "قرمز"
    assert pi.strip_count("۱") == ""


def test_norm_code_handles_excel_floats_and_persian_digits():
    assert pi.norm_code(1234.0) == "1234"
    assert pi.norm_code("۱۲۳۴") == "1234"
    assert pi.norm_code(" ab 12 ") == "AB12"


def test_is_featured():
    assert pi.is_featured(Path("x/main.jpg"))
    assert pi.is_featured(Path("x/IMG-Hero.png"))
    assert not pi.is_featured(Path("x/IMG_001.jpg"))


# --- color names -------------------------------------------------------------

def test_canonical_color():
    assert pi.canonical_color("navy-blue") == "navy"
    assert pi.canonical_color("off-white") == "white"
    assert pi.canonical_color("nescafe") == "coffee"
    assert pi.canonical_color("light-blue") == "light-blue"


def test_resolve_color_ignores_counts_and_matches_whole_words():
    cmap = {
        pi.norm_text("آبی"): "blue",
        pi.norm_text("آبی روشن"): "light-blue",
        pi.norm_text("نسکافه ای"): "nescafe",
        pi.norm_text("رز"): "rose",
    }
    assert pi.resolve_color("آبی ۱", cmap) == "blue"
    assert pi.resolve_color("آبی روشن 2", cmap) == "light-blue"
    assert pi.resolve_color("نسکافه‌ای ۲", cmap) == "nescafe"
    # "رز" must not match inside another word
    assert pi.resolve_color("گلرزی", cmap) is None
    assert pi.resolve_color("", cmap) is None


# --- color detection ---------------------------------------------------------

@pytest.mark.parametrize(
    "rgb, expected",
    [
        ((200, 30, 30), "red"),
        ((40, 90, 180), "blue"),
        ((20, 20, 20), "black"),
        ((128, 128, 128), "gray"),
    ],
)
def test_detect_color_ignores_white_background(rgb, expected):
    assert pi.detect_color(on_white(rgb)) == expected


def test_detect_color_respects_alpha_mask():
    img = on_white((200, 30, 30), size=200, box=100)
    alpha = np.zeros((200, 200), dtype=bool)
    alpha[50:150, 50:150] = True
    assert pi.detect_color(img, alpha) == "red"


def test_pick_color_exact_match_wins():
    allowed = [("dark-red", ref("dark-red")), ("red", ref("red"))]
    color, detected, _ = pi.pick_color(on_white((200, 30, 30)), None, allowed)
    assert (color, detected) == ("red", "red")


def test_pick_color_uses_nearest_when_not_allowed():
    # beige is detected, the product only has cream / brown -> cream
    allowed = [("brown", ref("brown")), ("cream", ref("cream"))]
    color, detected, _ = pi.pick_color(on_white((222, 211, 178)), None, allowed)
    assert detected == "beige"
    assert color == "cream"


def test_pick_color_chooses_closest_of_partial_matches():
    allowed = [("light-blue", ref("light-blue")), ("dark-blue", ref("dark-blue"))]
    color, detected, _ = pi.pick_color(on_white((30, 60, 125)), None, allowed)
    assert detected == "blue"
    assert color == "dark-blue"


def test_pick_color_without_allowed_returns_detected():
    color, detected, hexcode = pi.pick_color(on_white((200, 30, 30)), None, [])
    assert color == detected == "red"
    assert hexcode.startswith("#")


# --- compression -------------------------------------------------------------

def test_compress_leaves_file_that_fits_limit(tmp_path):
    rng = np.random.default_rng(0)
    arr = rng.integers(0, 256, (400, 400, 3), dtype=np.uint8)
    img = Image.fromarray(arr)
    dest = tmp_path / "noise.webp"
    limit = 40_000
    quality = pi.compress(img, dest, max_bytes=limit)
    assert dest.stat().st_size <= limit
    assert quality >= pi.QUALITY_FLOOR


def test_compress_small_image_uses_start_quality(tmp_path):
    dest = tmp_path / "flat.webp"
    assert pi.compress(solid((10, 20, 30)), dest) == pi.QUALITY_START


# --- end to end --------------------------------------------------------------

@pytest.fixture
def workspace(tmp_path, monkeypatch):
    photoroom = tmp_path / "photoroom"
    output = tmp_path / "final"
    for folder, name, color in [
        ("AB1", "main.png", (40, 90, 180)),     # blue, featured
        ("AB1", "b.png", (200, 30, 30)),         # red
        ("AB1", "c.png", (200, 30, 30)),         # red again -> 02
        ("CD2", "x.png", (200, 30, 30)),         # single color product -> blue
        ("ZZ9", "y.png", (200, 30, 30)),         # not in database
    ]:
        (photoroom / folder).mkdir(parents=True, exist_ok=True)
        on_white(color).save(photoroom / folder / name)

    db = tmp_path / "db.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "کد محصول", "رنگ"])
    ws.append([1, "AB1", "آبی ۱ | قرمز ۲"])
    ws.append([2, "CD2", "آبی"])
    ws.append([3, "EF3", "فیروزه‌ایی"])  # no mapping
    wb.save(db)

    cmap = tmp_path / "map.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Persian", "English"])
    ws.append(["آبی", "Blue"])
    ws.append(["قرمز", "Red"])
    wb.save(cmap)

    monkeypatch.setattr(pi, "PHOTOROOM_ROOT", photoroom)
    monkeypatch.setattr(pi, "OUTPUT_ROOT", output)
    monkeypatch.setattr(pi, "PRODUCT_DB", db)
    monkeypatch.setattr(pi, "COLOR_MAP", cmap)
    monkeypatch.setattr(pi, "COLOR_REPORT", tmp_path / "color.xlsx")
    return tmp_path


def names(folder: Path):
    return sorted(p.name for p in folder.iterdir() if p.suffix == ".png")


def test_process_renames_by_color_without_converting(workspace):
    pi.process(pi.find_photoroom_sources())
    out = workspace / "final"
    assert names(out / "AB1") == ["AB1-blue-main.png", "AB1-red-01.png", "AB1-red-02.png"]
    assert names(out / "CD2") == ["CD2-blue-01.png"]       # single color, no detection
    assert names(out / "ZZ9") == ["ZZ9-red-01.png"]        # detected only
    # copied unchanged
    src = Image.open(workspace / "photoroom" / "AB1" / "b.png")
    dst = Image.open(out / "AB1" / "AB1-red-01.png")
    assert src.size == dst.size and src.tobytes() == dst.tobytes()

    wb = openpyxl.load_workbook(workspace / "color.xlsx")
    missing = [row for row in wb["missing"].iter_rows(min_row=2, values_only=True)]
    problems = {(r[0], r[2]) for r in missing}
    assert ("EF3", "no entry in color mapping") in problems
    assert ("ZZ9", "code not found in database") in problems


def test_process_is_idempotent_and_recovers_deleted_outputs(workspace):
    pi.process(pi.find_photoroom_sources())
    out = workspace / "final"
    before = names(out / "AB1")
    pi.process(pi.find_photoroom_sources())
    assert names(out / "AB1") == before

    (out / "AB1" / "AB1-red-01.png").unlink()
    pi.process(pi.find_photoroom_sources())
    assert names(out / "AB1") == before


def test_manifest_is_saved_per_image(workspace, monkeypatch):
    calls = []
    real = pi.save_manifest
    monkeypatch.setattr(pi, "save_manifest", lambda m: (calls.append(len(m)), real(m)))
    pi.process(pi.find_photoroom_sources())
    assert max(calls) == 5 and calls[0] == 1


def test_corrupt_image_is_reported_not_fatal(workspace):
    bad = workspace / "photoroom" / "AB1" / "bad.png"
    bad.write_bytes(b"not an image")
    pi.process(pi.find_photoroom_sources())
    wb = openpyxl.load_workbook(workspace / "color.xlsx")
    problems = [r[2] for r in wb["missing"].iter_rows(min_row=2, values_only=True)]
    assert any(p.startswith("could not read image") for p in problems)
    assert len(names(workspace / "final" / "AB1")) == 3


def test_convert_handles_same_stem_collision(tmp_path, monkeypatch):
    root = tmp_path / "src"
    (root / "P1").mkdir(parents=True)
    solid((10, 20, 30)).save(root / "P1" / "a.jpg")
    solid((30, 20, 10)).save(root / "P1" / "a.png")
    monkeypatch.setattr(pi, "ROOT", root)
    monkeypatch.setattr(pi, "CONVERT_ROOT", tmp_path / "conv")
    out = pi.convert(pi.find_sources())
    assert len({p for p in out}) == 2
    assert all(p.exists() for p in out)
    assert not list((tmp_path / "conv").rglob("*.part.webp"))


def test_process_aborts_before_touching_images_when_database_missing(workspace, monkeypatch):
    monkeypatch.setattr(pi, "PRODUCT_DB", workspace / "nope.xlsx")
    with pytest.raises(SystemExit) as exc:
        pi.process(pi.find_photoroom_sources())
    assert "product database not found" in str(exc.value)
    assert not (workspace / "final").exists()


def test_locate_file_tolerates_small_name_differences(tmp_path):
    real = tmp_path / "Product Database - filled.xlsx"
    real.write_bytes(b"")
    found, _ = pi.locate_file(tmp_path / "Product Database .xlsx")
    assert found == real
    other = tmp_path / "Product Database v2.xlsx"
    other.write_bytes(b"")
    found, candidates = pi.locate_file(tmp_path / "Product Database .xlsx")
    assert found is None and len(candidates) == 2  # ambiguous: list them, don't guess


def test_slug_strips_accents_instead_of_dropping_letters():
    assert pi.slug("Café au lait / Nescafé") == "cafe-au-lait-nescafe"


def test_cafe_au_lait_keeps_its_name_and_has_a_reference():
    name = pi.canonical_color(pi.slug("Café au lait / Nescafé"))
    assert name == "cafe-au-lait"
    for color in ("walnut", "mocha", "cafe-au-lait"):
        assert pi.color_ref(color) is not None


def test_brown_object_maps_to_walnut_or_mocha_not_navy():
    allowed = [("navy", ref("navy")), ("mocha", ref("mocha"))]
    color, _, _ = pi.pick_color(on_white((111, 78, 55)), None, allowed)
    assert color == "mocha"


def test_colors_without_reference_are_reported(workspace, monkeypatch):
    wb = openpyxl.load_workbook(workspace / "db.xlsx")
    ws = wb.active
    ws.append([4, "GH4", "آبی | قرمز | زرشکی"])
    wb.save(workspace / "db.xlsx")
    cmap = openpyxl.load_workbook(workspace / "map.xlsx")
    cmap.active.append(["زرشکی", "Zorgh"])
    cmap.save(workspace / "map.xlsx")
    _, missing = pi.load_database()
    assert ("GH4", "zorgh", "no reference color: add it to color_refs in settings.json") in missing


# real colors measured on FA-1959 photos (median of the bag pixels)
FA_1959 = ["black", "cream", "cafe-au-lait", "pistachio", "emerald", "navy"]


@pytest.mark.parametrize(
    "rgb, expected",
    [
        ((32, 29, 27), "black"),
        ((37, 60, 38), "emerald"),        # dark green was matched to black before
        ((54, 59, 75), "navy"),           # navy bag was "detected black"
        ((72, 98, 72), "emerald"),        # green bag was "detected black"
        ((178, 146, 116), "cafe-au-lait"),
        ((129, 137, 105), "pistachio"),   # sage green was matched to cafe-au-lait
        ((219, 191, 154), "cream"),
        ((37, 40, 54), "navy"),
    ],
)
def test_fa_1959_real_colors(rgb, expected):
    allowed = [(n, ref(n)) for n in FA_1959]
    color, _, _ = pi.pick_color(on_white(rgb), None, allowed)
    assert color == expected


def test_pick_color_single_allowed_needs_no_reference():
    color, _, _ = pi.pick_color(on_white((200, 30, 30)), None, [("unknown-shade", None)])
    assert color == "unknown-shade"
